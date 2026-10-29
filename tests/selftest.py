"""resume-agent 自检脚本（不联网、不调用真实豆包 API）。

覆盖：
- 数据层增删改查、部署状态流转
- 简历模板两种填充模式（占位符 / 无占位符追加）
- Flask 全路由冒烟：设置、资料、项目、岗位分析（模拟豆包）、部署、简历生成与下载

运行：python tests/selftest.py
"""
from __future__ import annotations

import io
import os
import re
import shutil
import sys
import tempfile

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

PASS = 0
FAIL = 0


def check(name: str, cond: bool, detail: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✔ {name}")
    else:
        FAIL += 1
        print(f"  ✘ {name}  {detail}")


def make_placeholder_template(path: str) -> None:
    from docx import Document
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("{{NAME}}")
    p2 = doc.add_paragraph()
    p2.add_run("联系方式：{{PHONE}} / {{EMAIL}}")
    p3 = doc.add_paragraph()
    p3.add_run("个人简介：{{SUMMARY}}")
    p4 = doc.add_paragraph()
    p4.add_run("技能：{{SKILLS}}")
    doc.add_paragraph("项目经历")
    doc.add_paragraph("{{PROJECTS}}")
    doc.add_paragraph("实习经历")
    doc.add_paragraph("{{EXPERIENCE}}")
    doc.save(path)


def make_plain_template(path: str) -> None:
    from docx import Document
    doc = Document()
    doc.add_paragraph("张三简历")
    p = doc.add_paragraph()
    p.add_run("项目经历").bold = True
    doc.add_paragraph("个人信息简介……")
    doc.save(path)


def read_docx_text(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    parts.append(p.text)
    return "\n".join(parts)


class FakeClient:
    """模拟豆包客户端：按 prompt 特征返回预设 JSON。"""
    def chat_json(self, messages, **kwargs):
        user = messages[-1]["content"]
        if "候选项目" in user and "id=" in user:  # match_projects
            pid = re.search(r'id="([^"]+)"', user).group(1)
            return {
                "results": [{"id": pid, "score": 88, "verdict": "high",
                             "reason": "技术栈命中岗位要求", "highlight_points": ["数据处理经验", "Flask 部署"],
                             "gap": ""}],
                "best_project_id": pid, "has_good_match": True,
                "missing_skills": [], "summary": "测试项目高度匹配该岗位。",
            }
        if "改写成简历条目" in user:  # rewrite_project_for_job
            return {"title": "测试项目 · 核心开发",
                    "bullets": ["要点1：独立完成核心模块开发", "要点2：接口耗时降低 40%"],
                    "deploy_line": ""}
        if "结构化" in user and "技术栈" in user:  # analyze_project_text
            url = re.search(r"https?://\S+", user)
            return {"name": "AI导入项目", "role": "开发", "tech_stack": ["Python"],
                    "description": "AI 解析的项目", "highlights": ["亮点A"],
                    "source_url": url.group(0) if url else ""}
        if "候选开源项目" in user:  # rank_oss_projects
            return {"picks": []}
        if "判断它与目标岗位的适配度" in user:  # analyze_experience
            return {"items": [
                {"raw": "某事务所 · 数据分析实习生：负责订单数据清洗，处理10万条订单数据",
                 "include": True,
                 "reason": "数据处理与岗位技能高度匹配",
                 "rewritten": "负责订单数据清洗，处理 10 万条订单数据，产出可视化报表\n协助优化数据处理流程，报表产出时间缩短 30%"},
                {"raw": "某教育机构 · 助教：负责批改作业、维持课堂秩序",
                 "include": False,
                 "reason": "与后端开发/数据分析岗位无关",
                 "rewritten": ""},
            ]}
        # analyze_job
        return {"position": "Python开发工程师", "company": "测试公司",
                "key_skills": [{"skill": "Python", "weight": 0.9}, {"skill": "Flask", "weight": 0.8},
                               {"skill": "Pandas", "weight": 0.6}],
                "responsibilities": ["负责后端开发与部署"], "requirements": ["本科及以上"],
                "education_requirement": "本科", "salary": "8-12k",
                "keywords": ["python", "flask", "数据分析"]}


def main():
    print("== 准备隔离数据目录 ==")
    tmp = tempfile.mkdtemp(prefix="resume_agent_selftest_")
    try:
        from resume_agent import store
        for d in ("", "templates", "resumes", "cache"):
            os.makedirs(os.path.join(tmp, d), exist_ok=True)
        store.DATA_DIR = tmp
        store.TEMPLATE_DIR = os.path.join(tmp, "templates")
        store.RESUME_DIR = os.path.join(tmp, "resumes")
        store.CACHE_DIR = os.path.join(tmp, "cache")
        store.SETTINGS_FILE = os.path.join(tmp, "settings.json")
        store.PROFILE_FILE = os.path.join(tmp, "profile.json")
        store.PROJECTS_FILE = os.path.join(tmp, "projects.json")
        store.ANALYSIS_FILE = os.path.join(tmp, "cache", "last_analysis.json")
        store.REPORT_FILE = os.path.join(tmp, "cache", "last_report.json")
        store.OSS_CACHE_FILE = os.path.join(tmp, "cache", "oss_candidates.json")

        print("== 数据层测试 ==")
        store.save_settings({"llm_provider": "doubao", "doubao_api_key": "test-key", "doubao_model": "m"})
        s = store.load_settings()
        cfg = store.get_llm_config()
        check("设置保存/读取", s["llm_provider"] == "doubao" and cfg["api_key"] == "test-key")
        check("默认服务商为豆包", cfg["provider"] == "doubao" and "ark.cn-beijing" in cfg["base_url"])

        # 旧版配置（ark_ 前缀）迁移
        store.save_json(store.SETTINGS_FILE,
                        {"ark_api_key": "old-key", "ark_base_url": "https://old", "ark_model": "old-model"})
        s2 = store.load_settings()
        check("旧版配置迁移", s2["doubao_api_key"] == "old-key" and s2["doubao_model"] == "old-model")
        store.save_settings({"llm_provider": "doubao", "doubao_api_key": "test-key", "doubao_model": "m"})

        # 服务商预设
        cfg_ds = store.get_llm_config({"llm_provider": "deepseek"})
        check("deepseek预设", cfg_ds["base_url"] == "https://api.deepseek.com"
              and cfg_ds["model"] == "deepseek-chat" and cfg_ds["api_key"] == "")
        cfg_oa = store.get_llm_config({"llm_provider": "openai"})
        check("openai预设", cfg_oa["base_url"] == "https://api.openai.com/v1" and cfg_oa["model"] == "gpt-4o-mini")
        cfg_cc = store.get_llm_config({"llm_provider": "custom", "custom_api_key": "ck",
                                       "custom_base_url": "https://my/v1", "custom_model": "my-model"})
        check("自定义服务商", cfg_cc["base_url"] == "https://my/v1" and cfg_cc["model"] == "my-model"
              and cfg_cc["api_key"] == "ck")
        cfg_ce = store.get_llm_config({"llm_provider": "custom"})
        check("自定义缺省为空", cfg_ce["base_url"] == "" and cfg_ce["model"] == "")

        # 无 Key 时 LLMClient 构造报错
        from resume_agent.llm_client import LLMClient, LLMError
        try:
            LLMClient("", "https://api.deepseek.com", "deepseek-chat")
            check("无Key构造报错", False)
        except LLMError:
            check("无Key构造报错", True)
        p1 = store.add_project({"name": "测试项目", "role": "核心开发",
                                "tech_stack": ["Python", "Flask"],
                                "description": "数据处理平台", "highlights": "接口快\n并发高"})
        check("项目入库", len(store.list_projects()) == 1)
        check("亮点字符串转列表", isinstance(p1["highlights"], list) and len(p1["highlights"]) == 2)
        store.mark_deployed(p1["id"], "http://127.0.0.1:5000")
        check("部署状态更新", store.get_project(p1["id"])["deployed"] is True)

        print("== 简历模板测试 ==")
        from resume_agent import resume_builder
        tpl_ph = os.path.join(tmp, "tpl_placeholder.docx")
        make_placeholder_template(tpl_ph)
        out1 = os.path.join(tmp, "out_placeholder.docx")
        resume_builder.fill_template(tpl_ph, {
            "single": {"name": "张三", "phone": "13800000000", "email": "a@b.com",
                       "summary": "熟悉Python", "skills": "Python, Flask"},
            "projects": [{"title": "测试项目 · 核心开发",
                          "bullets": ["要点1", "要点2"],
                          "deploy_line": "已部署：http://127.0.0.1:5000"}],
            "experience": "某事务所实习",
        }, out1)
        t1 = read_docx_text(out1)
        check("占位符模式-姓名替换", "张三" in t1)
        check("占位符模式-技能替换", "Python, Flask" in t1)
        check("占位符模式-项目标题插入", "测试项目 · 核心开发" in t1)
        check("占位符模式-要点插入", "要点1" in t1 and "要点2" in t1)
        check("占位符模式-部署行插入", "已部署：http://127.0.0.1:5000" in t1)
        check("占位符模式-实习经历插入", "某事务所实习" in t1)

        tpl_plain = os.path.join(tmp, "tpl_plain.docx")
        make_plain_template(tpl_plain)
        out2 = os.path.join(tmp, "out_plain.docx")
        resume_builder.fill_template(tpl_plain, {
            "single": {}, "projects": [{"title": "开源项目实践", "bullets": ["完成部署"], "deploy_line": ""}],
            "experience": "",
        }, out2)
        t2 = read_docx_text(out2)
        check("追加模式-项目标题", "开源项目实践" in t2)
        check("追加模式-要点", "完成部署" in t2)
        # 标题后插入（顺序：标题行应在项目条目前）
        check("追加模式-插在项目经历标题后",
              t2.find("项目经历") < t2.find("开源项目实践"))

        print("== Flask 路由测试（模拟豆包） ==")
        import app as app_mod
        app_mod.get_client = lambda: FakeClient()
        client = app_mod.app.test_client()

        for path in ("/", "/settings", "/templates", "/profile", "/projects", "/analyze", "/recommend", "/resume"):
            r = client.get(path)
            check(f"GET {path} -> {r.status_code}", r.status_code == 200)

        r = client.post("/profile", data={"name": "张三", "phone": "13800000000",
                                          "school": "测试大学", "major": "数据科学"})
        check("POST /profile 重定向", r.status_code in (302, 303))

        r = client.post("/projects", data={"name": "测试项目", "role": "核心开发",
                                           "tech_stack": "Python, Flask",
                                           "description": "数据处理平台", "highlights": "接口快"})
        check("POST /projects 添加", r.status_code in (302, 303))
        pid = store.list_projects()[0]["id"]

        r = client.get("/settings")
        check("GET /settings 渲染三家服务商", r.status_code == 200
              and "DeepSeek" in r.get_data(as_text=True) and "OpenAI" in r.get_data(as_text=True))
        r = client.post("/settings", data={"llm_provider": "deepseek",
                                           "deepseek_api_key": "ds-key",
                                           "openai_api_key": "oa-key"})
        check("POST /settings 切换服务商", r.status_code in (302, 303))
        cfg_saved = store.get_llm_config()
        check("切换后生效", cfg_saved["provider"] == "deepseek" and cfg_saved["api_key"] == "ds-key"
              and store.load_settings()["openai_api_key"] == "oa-key")
        r = client.get("/")
        check("首页显示当前模型", r.status_code == 200
              and "DeepSeek" in r.get_data(as_text=True) and "deepseek-chat" in r.get_data(as_text=True))
        r = client.post("/settings", data={"llm_provider": "doubao", "doubao_api_key": "test-key"})
        check("切回豆包", r.status_code in (302, 303) and store.get_llm_config()["api_key"] == "test-key")

        # 项目编辑
        r = client.post(f"/projects/edit/{pid}", data={"name": "测试项目改", "role": "全栈",
                                                       "tech_stack": "Python, Django",
                                                       "description": "改过的描述", "highlights": "新亮点A\n新亮点B"})
        check("POST /projects/edit 保存", r.status_code in (302, 303)
              and store.get_project(pid)["name"] == "测试项目改"
              and store.get_project(pid)["highlights"] == ["新亮点A", "新亮点B"])
        r = client.get(f"/projects?edit={pid}")
        check("编辑表单回填", r.status_code == 200 and "测试项目改" in r.get_data(as_text=True)
              and "保存修改" in r.get_data(as_text=True))
        r = client.post(f"/projects/edit/{pid}", data={"name": "测试项目", "role": "核心开发",
                                                       "tech_stack": "Python, Flask",
                                                       "description": "数据处理平台", "highlights": "接口快"})
        check("改回原名", r.status_code in (302, 303) and store.get_project(pid)["name"] == "测试项目")

        # 上传模板
        with open(tpl_ph, "rb") as f:
            r = client.post("/template/upload",
                            data={"template": (io.BytesIO(f.read()), "我的模板.docx")},
                            content_type="multipart/form-data")
        check("模板上传", r.status_code in (302, 303) and len(store.list_uploaded_templates()) == 1)
        tpl_name = store.list_uploaded_templates()[0]
        r = client.get("/templates")
        check("模板专属页", r.status_code == 200 and tpl_name in r.get_data(as_text=True))
        r = client.post(f"/template/delete/{tpl_name}")
        check("模板删除", r.status_code in (302, 303) and len(store.list_uploaded_templates()) == 0)
        with open(tpl_ph, "rb") as f:
            r = client.post("/template/upload",
                            data={"template": (io.BytesIO(f.read()), "我的模板.docx")},
                            content_type="multipart/form-data")
        tpl_name = store.list_uploaded_templates()[0]

        # 无 Key 时的错误兜底（清掉 key 模拟）
        store.save_settings({"llm_provider": "doubao", "doubao_api_key": ""})
        r = client.post("/analyze", data={"job_text": "这是岗位信息，要求Python开发工程师，熟悉Flask、爬虫与数据分析，负责后端接口开发与部署。" * 2})
        check("无Key-岗位分析有错误提示", r.status_code in (302, 303))
        r = client.post("/resume/generate", data={"template": tpl_name, "project_ids": [pid]})
        check("无Key-简历生成有错误提示", r.status_code in (302, 303))
        store.save_settings({"llm_provider": "doubao", "doubao_api_key": "test-key"})

        # 岗位分析全流程（模拟豆包）
        job_text = "岗位：Python开发工程师。职责：后端接口开发、爬虫、数据分析。要求：熟悉Python/Flask/Pandas，本科以上。"
        r = client.post("/analyze", data={"job_text": job_text})
        check("POST /analyze 分析成功", r.status_code in (302, 303))
        analysis = store.load_analysis()
        report = store.load_report()
        check("分析结果已缓存", analysis and analysis["position"] == "Python开发工程师")
        check("匹配报告已缓存", report and report["has_good_match"] is True)
        r = client.get("/analyze")
        check("GET /analyze 展示匹配卡片", r.status_code == 200 and "88" in r.get_data(as_text=True))

        # 部署确认
        r = client.post(f"/projects/{pid}/deployed", data={"deploy_url": "http://127.0.0.1:5000"})
        check("部署确认", r.status_code in (302, 303) and store.get_project(pid)["deployed"] is True)
        r = client.get(f"/deploy/{pid}")
        check("GET /deploy 部署页", r.status_code == 200)

        # 生成简历全流程（模拟豆包改写）
        r = client.post("/resume/generate", data={"template": tpl_name, "project_ids": [pid]})
        check("POST /resume/generate 生成", r.status_code in (302, 303))
        generated = store.list_generated_resumes()
        check("简历文件已产出", len(generated) >= 1)
        if generated:
            fname = generated[0]
            r = client.get(f"/resume/download/{fname}")
            check("简历可下载", r.status_code == 200 and "attachment" in r.headers.get("Content-Disposition", ""))
            body = read_docx_text(os.path.join(store.RESUME_DIR, fname))
            check("简历内容包含改写条目", "要点1" in body and "张三" in body)

        # 工作经历：按岗位适配度过滤 + 数据化改写
        r = client.post("/profile", data={"name": "张三", "phone": "13800000000",
                                          "school": "测试大学", "major": "数据科学",
                                          "experience": "某事务所 · 数据分析实习生：负责订单数据清洗，处理10万条订单数据\n某教育机构 · 助教：负责批改作业、维持课堂秩序"})
        check("保存含工作经历的资料", r.status_code in (302, 303))
        r = client.post("/resume/generate", data={"template": tpl_name, "project_ids": [pid],
                                                  "filter_experience": "on"})
        check("过滤模式生成成功", r.status_code in (302, 303))
        f2 = store.list_generated_resumes()[0]
        body2 = read_docx_text(os.path.join(store.RESUME_DIR, f2))
        check("写入适配且改写的工作经历", "10 万条订单数据" in body2 and "报表产出时间缩短 30%" in body2)
        check("不适配经历被过滤", "助教" not in body2)
        r = client.get("/resume")
        check("简历页展示适配分析结果", r.status_code == 200
              and "写入简历" in r.get_data(as_text=True) and "不写入" in r.get_data(as_text=True))
        exp_cache = store.load_experience_analysis()
        check("分析结果已缓存", exp_cache is not None and len(exp_cache.get("items", [])) == 2)
        # 关闭过滤：按原文全部写入
        r = client.post("/resume/generate", data={"template": tpl_name, "project_ids": [pid]})
        check("原文模式生成成功", r.status_code in (302, 303))
        f3 = store.list_generated_resumes()[0]
        body3 = read_docx_text(os.path.join(store.RESUME_DIR, f3))
        check("原文模式写入全部经历", "助教" in body3 and "10万条" in body3)

        # AI 识别回填（不直接入库，供前端表单回填）
        r = client.post("/projects/parse", data={"text": "我做的项目，用了Python和爬虫，链接 https://github.com/x/y"})
        body = r.get_json() if r.is_json else None
        check("AI识别返回JSON", r.status_code == 200 and body and body.get("ok") is True)
        check("AI识别提取字段", bool(body) and body["project"]["name"] == "AI导入项目"
              and body["project"]["source_url"] == "https://github.com/x/y")
        check("AI识别未直接入库", not any(p["name"] == "AI导入项目" for p in store.list_projects()))
        r = client.post("/projects/parse", data={"text": ""})
        check("AI识别空文本报错", r.status_code == 400 and r.get_json().get("ok") is False)

        # 岗位分析历史
        history = store.list_history()
        check("分析后自动存历史", len(history) >= 1
              and any(h["analysis"]["position"] == "Python开发工程师" for h in history))
        hid = history[0]["id"]
        r = client.post(f"/analyze/delete/{hid}")
        check("历史删除", r.status_code in (302, 303)
              and not any(h["id"] == hid for h in store.list_history()))
        # 再分析一次，用于后续测试
        r = client.post("/analyze", data={"job_text": job_text})
        hid = store.list_history()[0]["id"]
        r = client.get(f"/analyze?history={hid}")
        check("载入历史记录", r.status_code == 200 and "已载入历史记录" in r.get_data(as_text=True))

        # 开源搜索关键词：精简 + 手动
        from resume_agent import oss_client
        q1 = oss_client.build_query({"key_skills": [{"skill": "Flask", "weight": 0.9},
                                                    {"skill": "Pandas", "weight": 0.6}], "keywords": []})
        check("关键词自动精简", q1 == "Flask")
        q2 = oss_client.build_query({"key_skills": [{"skill": "Python", "weight": 0.9},
                                                    {"skill": "爬虫", "weight": 0.8}], "keywords": []})
        check("泛词补第二词", q2 == "Python 爬虫")
        q3 = oss_client.build_query({}, "flask 爬虫")
        check("手动关键词优先", q3 == "flask 爬虫")
        # 模拟搜索（不联网）：固定候选 + 验证手动关键词透传
        fake_cands = [{"full_name": "a/flask-app", "html_url": "https://x/a", "description": "d", "stars": 99,
                       "language": "Python", "topics": [], "platform": "github"},
                      {"full_name": "b/crawler", "html_url": "https://x/b", "description": "d", "stars": 88,
                       "language": "Python", "topics": [], "platform": "github"}]
        app_mod.oss_client.search_any = lambda analysis, settings, manual="", order=("github", "gitee"): {
            "platform": "github", "query": manual or oss_client.build_query(analysis),
            "candidates": fake_cands, "errors": []}
        r = client.post("/recommend/search", data={"keywords": "flask 爬虫"})
        check("手动关键词搜索", r.status_code in (302, 303))
        cached = store.load_oss_candidates()
        check("候选已缓存且关键词保留", isinstance(cached, dict) and cached.get("query") == "flask 爬虫"
              and len(cached.get("candidates", [])) >= 1)
        r = client.get("/recommend")
        check("推荐页显示可编辑关键词", r.status_code == 200 and "flask 爬虫" in r.get_data(as_text=True))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print(f"\n结果：通过 {PASS} 项，失败 {FAIL} 项")
    sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
