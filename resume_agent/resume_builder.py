"""简历生成：把个人资料 + 岗位匹配结果 + 改写后的项目经历，填充进用户上传的 Word 模板。

支持两种模式（自动检测）：
1. 占位符模式：模板里含 {{NAME}} {{PROJECTS}} 等占位符，逐一替换/插入。
2. 追加模式：模板无占位符，自动定位「项目经历」标题并在其后插入条目；找不到则追加到文末。

填充时通过复制模板段落 XML 保留原有字体、字号、样式，中文显示不受影响。
"""
from __future__ import annotations

import os
import re
from copy import deepcopy
from datetime import datetime

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

# 单值占位符 -> 简历数据键
SINGLE_PLACEHOLDERS = {
    "NAME": "name",
    "PHONE": "phone",
    "EMAIL": "email",
    "CITY": "city",
    "SCHOOL": "school",
    "MAJOR": "major",
    "DEGREE": "degree",
    "GRADUATION": "graduation",
    "TARGET_POSITION": "target_position",
    "SUMMARY": "summary",
    "SKILLS": "skills",
    "CERTIFICATES": "certificates",
    "LANGUAGES": "languages",
}
# 多段占位符 -> 简历数据键（每个值内部按行拆分）
MULTI_PLACEHOLDERS = {
    "PROJECTS": "projects",
    "EXPERIENCE": "experience",
}
ALL_PLACEHOLDERS = set(SINGLE_PLACEHOLDERS) | set(MULTI_PLACEHOLDERS)
TOKEN_RE = re.compile(r"\{\{\s*([A-Z_]+)\s*\}\}")


class ResumeBuildError(Exception):
    pass


def _set_paragraph_text(p_element, text: str, bold: bool = False) -> None:
    """把段落 XML 的文本设为 text，保留首个 run 的字体格式。"""
    runs = p_element.findall(qn("w:r"))
    if not runs:
        r = OxmlElement("w:r")
        p_element.append(r)
        runs = [r]
    first = runs[0]
    for r in runs[1:]:
        p_element.remove(r)
    for t in first.findall(qn("w:t")):
        first.remove(t)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    first.append(t)
    if bold:
        rPr = first.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            first.insert(0, rPr)
        if rPr.find(qn("w:b")) is None:
            rPr.append(OxmlElement("w:b"))


def _iter_all_paragraphs(doc):
    """遍历正文段落 + 所有表格单元格内的段落。"""
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def _find_placeholder_paragraphs(doc):
    """返回 {placeholder: [paragraph, ...]}，按文档顺序。"""
    found = {}
    for p in _iter_all_paragraphs(doc):
        for token in TOKEN_RE.findall(p.text or ""):
            if token in ALL_PLACEHOLDERS:
                found.setdefault(token, []).append(p)
    return found


def _replace_single_in_paragraph(p, data: dict) -> bool:
    """把段落内出现的单值占位符替换为文本，返回是否发生过替换。"""
    text = p.text or ""
    if not text:
        return False

    def repl(m):
        token = m.group(1)
        if token not in SINGLE_PLACEHOLDERS:
            return m.group(0)
        value = data.get(SINGLE_PLACEHOLDERS[token]) or ""
        return str(value).replace("\n", " ")

    new_text = TOKEN_RE.sub(repl, text)
    if new_text == text:
        return False
    _set_paragraph_text(p._p, new_text)
    return True


def _insert_lines_before(anchor_par, lines) -> None:
    """在 anchor 段落之前插入若干行（deepcopy 其 XML 保留格式）。"""
    for item in reversed(lines):
        text, bold = item
        new_el = deepcopy(anchor_par._p)
        anchor_par._p.addprevious(new_el)
        _set_paragraph_text(new_el, text, bold=bold)


def _remove_paragraph(p) -> None:
    p._p.getparent().remove(p._p)


def _project_lines(project_entries: list) -> list:
    """把项目条目转换为 [(文本, 是否加粗), ...] 行序列。"""
    lines = []
    for entry in project_entries:
        title = entry.get("title") or entry.get("name") or "项目"
        lines.append((title, True))
        for b in entry.get("bullets") or []:
            lines.append((f"• {b}", False))
        deploy_line = entry.get("deploy_line") or ""
        if deploy_line:
            lines.append((deploy_line, False))
    return lines


def _experience_lines(experience) -> list:
    """工作/实习经历转行序列。支持字符串（按行拆分）或已处理好的行列表。"""
    out = []
    if isinstance(experience, list):
        for line in experience:
            line = str(line).strip()
            if line:
                out.append((line, False))
        return out
    for line in str(experience or "").split("\n"):
        line = line.strip()
        if not line:
            continue
        out.append((line, False))
    return out


def _append_at_document_end(doc, lines, body_template_p=None) -> None:
    """追加到文末：若没有项目标题则先加一行加粗标题。"""
    anchor = doc.paragraphs[-1]._p if doc.paragraphs else None
    if anchor is None:
        # 文档为空：直接用 python-docx 新建
        p = doc.add_paragraph()
        run = p.add_run("项目经历")
        run.bold = True
        _set_run_cn_font(run)
        anchor = p._p
    else:
        template_el = body_template_p._p if body_template_p else anchor
        style_el = deepcopy(template_el)
        anchor.addnext(style_el)
        _set_paragraph_text(style_el, "项目经历", bold=True)
        anchor = style_el
    for text, bold in lines:
        new_el = deepcopy(anchor)
        anchor.addnext(new_el)
        _set_paragraph_text(new_el, text, bold=bold)
        anchor = new_el


def _set_run_cn_font(run, name: str = "微软雅黑") -> None:
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def _find_project_heading(doc):
    """找包含「项目」的标题段（Heading 样式或加粗段），供追加模式定位。"""
    for p in doc.paragraphs:
        if not (p.text or "").strip():
            continue
        if "项目" not in p.text:
            continue
        is_heading = p.style.name.startswith("Heading") or "标题" in p.style.name
        is_bold = bool(p.runs and p.runs[0].bold)
        if is_heading or is_bold:
            return p
    return None


def _first_body_paragraph(doc):
    for p in doc.paragraphs:
        if (p.text or "").strip() and not p.style.name.startswith("Heading"):
            return p
    return None


def fill_template(template_path: str, resume_data: dict, out_path: str) -> str:
    """填充模板并输出 docx，返回输出路径。"""
    if not os.path.exists(template_path):
        raise ResumeBuildError("找不到模板文件，请先在「设置」页上传 Word 模板")

    doc = Document(template_path)
    single_data = resume_data.get("single", {})
    project_entries = resume_data.get("projects", [])
    experience_text = resume_data.get("experience", "")

    found = _find_placeholder_paragraphs(doc)
    if found:
        # ---------- 占位符模式 ----------
        for p in _iter_all_paragraphs(doc):
            _replace_single_in_paragraph(p, single_data)

        for token, paras in found.items():
            if token in SINGLE_PLACEHOLDERS:
                continue  # 上面已处理
            key = MULTI_PLACEHOLDERS[token]
            if token == "PROJECTS":
                lines = _project_lines(project_entries)
            else:
                lines = _experience_lines(experience_text)
            if not lines:
                lines = [("暂无内容", False)]
            for anchor in paras:
                _insert_lines_before(anchor, lines)
            for anchor in paras:
                _remove_paragraph(anchor)
    else:
        # ---------- 追加模式 ----------
        lines = _project_lines(project_entries)
        exp_lines = _experience_lines(experience_text)
        heading = _find_project_heading(doc)
        body_t = _first_body_paragraph(doc)
        if heading is not None:
            anchor = heading._p
            template_el = body_t._p if body_t else heading._p
            for text, bold in lines:
                new_el = deepcopy(template_el)
                anchor.addnext(new_el)
                _set_paragraph_text(new_el, text, bold=bold)
                anchor = new_el
            if exp_lines:
                for text, bold in exp_lines:
                    new_el = deepcopy(template_el)
                    anchor.addnext(new_el)
                    _set_paragraph_text(new_el, text, bold=bold)
                    anchor = new_el
        else:
            all_lines = lines + exp_lines
            _append_at_document_end(doc, all_lines, body_template_p=body_t)

    doc.save(out_path)
    return out_path


def build_resume(profile: dict, project_entries: list, template_path: str,
                 position: str = "", out_dir: str = "", experience_lines: list | None = None) -> str:
    """组装简历数据并调用填充。

    project_entries 为 [{title, bullets, deploy_line}]；
    experience_lines 为已按岗位筛选/改写后的工作经历行列表（None 时使用个人资料原文）。
    """
    if not os.path.exists(template_path):
        raise ResumeBuildError("尚未上传简历模板，请先到「设置」页上传 .docx 模板")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    pos_slug = "".join(c for c in (position or "求职") if c not in r'\/:*?"<>|')[:20]
    out_name = f"简历_{pos_slug}_{ts}.docx"
    out_path = os.path.join(out_dir, out_name)

    single = {key: (profile.get(key) or "") for key in SINGLE_PLACEHOLDERS.values()}
    resume_data = {
        "single": single,
        "projects": project_entries,
        "experience": experience_lines if experience_lines is not None else profile.get("experience", ""),
    }
    fill_template(template_path, resume_data, out_path)
    return out_path
